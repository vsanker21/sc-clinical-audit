#!/usr/bin/env python3
"""Generate Archetype A manuscript (main + supplementary) from archetype_a artifacts."""

from __future__ import annotations

import argparse
import json
import shutil
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]

GB_ZENODO = "10.5281/zenodo.20610218"
GB_GITHUB = "https://github.com/vsanker21/scrna-metastasis-confound-audit"
AUTHOR_LINE = "Sanker, V. (TN Omics)"
CORRESPONDENCE = "[correspondence email to be completed]"


def _load_json(p: Path) -> Dict[str, Any]:
    return json.loads(p.read_text(encoding="utf-8"))


def _style(doc: Document) -> None:
    n = doc.styles["Normal"]
    n.font.name = "Times New Roman"
    n.font.size = Pt(12)


def _p(doc: Document, text: str, *, bold: bool = False, italic: bool = False) -> None:
    para = doc.add_paragraph()
    r = para.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)


def _h(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    h.runs[0].font.color.rgb = RGBColor(0, 0, 0)


def _fig(doc: Document, path: Path, legend: str, fig_id: str) -> None:
    if path.is_file():
        doc.add_picture(str(path), width=Inches(6.2))
    _p(doc, f"{fig_id}. {legend}", italic=True)


def _table_from_df(doc: Document, df: pd.DataFrame, title: str, max_rows: int = 25) -> None:
    _p(doc, title, bold=True)
    sub = df.head(max_rows)
    tbl = doc.add_table(rows=1, cols=len(sub.columns))
    tbl.style = "Table Grid"
    for j, c in enumerate(sub.columns):
        tbl.rows[0].cells[j].text = str(c)
    for _, row in sub.iterrows():
        cells = tbl.add_row().cells
        for j, c in enumerate(sub.columns):
            cells[j].text = str(row[c])[:80]
    doc.add_paragraph()


def _headline(dr: Path) -> Dict[str, Any]:
    p = dr / "artifacts/scientific_analyses/archetype_a/archetype_a_report.json"
    if not p.is_file():
        raise FileNotFoundError(f"Run archetype A study first: {p}")
    return _load_json(p)


def build_main(dr: Path, out: Path, fig_dir: Path) -> Path:
    report = _headline(dr)
    hc = report.get("headline_conclusions", {})
    reg = report.get("registry_summary", {})
    tcga = report.get("tcga_bulk_positive_control", {})

    doc = Document()
    _style(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run(
        "scClinicalAudit: a field benchmark for single-cell clinical-label prediction "
        "under indication, portal, and provenance confounding"
    )
    tr.bold = True
    tr.font.size = Pt(14)
    tr.font.name = "Arial"

    _p(doc, f"Authors: {AUTHOR_LINE}")
    _p(doc, f"Correspondence: {CORRESPONDENCE}")

    _h(doc, "Abstract", 1)
    _p(
        doc,
        "Background. Public single-cell RNA atlases are increasingly reused to predict clinical states such as "
        "metastasis, yet systematic audits for cancer-type, portal, and label-provenance confounding remain rare. "
        "We introduce scClinicalAudit, a reusable evaluation framework with four pre-specified estimands: confound "
        "fraction, transportability, signature graveyard survival, and bulk-RNA positive controls."
    )
    _p(
        doc,
        f"Results. Across a multi-drive registry of {reg.get('n_unique_specimens', 'n/a')} unique specimens "
        f"({reg.get('n_supervised', 'n/a')} supervised), apparent pan-cohort scRNA metastasis ranking on the "
        f"canonical specimen-authority cohort reached OOF AUROC {hc.get('canonical_task_auroc', 0):.3f} while a "
        f"combined metadata ceiling reached {hc.get('canonical_metadata_ceiling', 0):.3f} (confound fraction "
        f"{hc.get('canonical_confound_fraction', 0):.2f}). Of {hc.get('signatures_tested', 0)} curated proxy "
        f"programs, {hc.get('signatures_surviving_correction', 0)} survived strict within-indication correction "
        f"(corrected AUROC ≥0.60; Δ<0.10). A balanced TCGA-BRCA bulk RNA positive control retained OOF AUROC "
        f"{hc.get('tcga_bulk_positive_control', 0):.3f}, whereas a THCA null arm did not (OOF ≈0.47)."
    )
    _p(
        doc,
        "Conclusions. scClinicalAudit discriminates artifactual scRNA supervision from detectable within-indication "
        "bulk signal and provides a pilot field benchmark—not a metastasis-discovery claim. This work generalizes "
        f"a companion single-cohort audit (Genome Biology submission; Zenodo {GB_ZENODO})."
    )

    _h(doc, "Introduction", 1)
    _p(
        doc,
        "Computational reuse of public single-cell tumor-microenvironment (TME) atlases for supervised clinical "
        "labels has outpaced standardized evaluation. Models can rank specimens using indication tags, data portal, "
        "or label-provenance structure rather than metastasis biology. Donor-grouped cross-validation controls donor "
        "leakage but not indication-level confounding."
    )
    _p(
        doc,
        "We present scClinicalAudit Archetype A: a field benchmark that (i) builds a multi-drive specimen registry, "
        "(ii) decomposes apparent AUROC into metadata-explained and residual components, (iii) quantifies "
        "transportability across indications and portals, (iv) re-evaluates curated proxy signatures under "
        "within-indication correction, and (v) anchors interpretation with balanced bulk TCGA positive and null "
        "controls. This paper is a methods and benchmark extension—not a claim of metastasis biomarker discovery."
    )
    _p(
        doc,
        f"Relation to companion audit. A deep single-cohort confound audit of the same canonical labels "
        f"(142 supervised; 33 metastatic-positive) is reported separately (Zenodo {GB_ZENODO}; "
        f"{GB_GITHUB}), including calibration failure, block-permutation nulls, and HTAPP composition warnings. "
        "Here we generalize estimands across integrated packs and publish the reusable software."
    )

    _h(doc, "Results", 1)
    _h(doc, "Multi-drive specimen registry", 2)
    _p(
        doc,
        f"The registry deduplicated {reg.get('n_unique_specimens')} graph-linked specimens across C/D/E/F assets, "
        f"with {reg.get('n_supervised')} supervised rows under specimen-authority or pack-specific label policies. "
        "Leakage modes (specimen-grounded, portal–indication coupled, indication-saturated, inferred) are tabulated "
        "in Supplementary Table S1."
    )
    _fig(
        doc,
        fig_dir / "FigA4_leakage_taxonomy.png",
        "Registry breadth and leakage taxonomy across multi-drive manifests.",
        "Figure 1",
    )

    _h(doc, "Confound decomposition spectrum (E1)", 2)
    _p(
        doc,
        "For each integrated pack we estimated task OOF AUROC, adversarial metadata ceilings, and confound fraction "
        "(fraction of AUROC above chance explained by metadata). Canonical specimen-authority confound fraction was "
        f"{hc.get('canonical_confound_fraction', 0):.2f}. Headline scale verification via f142 D-graph symlink "
        f"parity yielded mean fold AUROC {hc.get('headline_scale_task_auroc', 0):.3f} with gate pass "
        f"{hc.get('headline_scale_parity_gate_pass')}."
    )
    _fig(
        doc,
        fig_dir / "FigA1_decomposition_spectrum.png",
        "Stacked decomposition of apparent AUROC above chance across integrated packs.",
        "Figure 2",
    )
    _fig(
        doc,
        fig_dir / "FigA6_confound_ceiling.png",
        "Adversarial metadata baselines versus TME block logistic regression on the canonical cohort.",
        "Figure 3",
    )

    _h(doc, "Transportability matrices (E2)", 2)
    _p(
        doc,
        "Train-stratum → test-stratum block-logistic models quantify whether apparent performance transports across "
        "indications. Off-diagonal degradation indicates indication-specific structure rather than universal biology. "
        "Portal-stratified transport is shown in Supplementary Figure S1."
    )
    _fig(
        doc,
        fig_dir / "FigA2_transportability_matrix.png",
        "Indication × indication transportability matrix (canonical cohort).",
        "Figure 4",
    )

    _h(doc, "Signature graveyard (E3)", 2)
    _p(
        doc,
        f"We evaluated {hc.get('signatures_tested', 0)} curated proxy programs (literature gene sets, TME composition "
        f"proxies, internal DE/GSEA panels). Primary survival used strict criteria: within-indication corrected OOF "
        f"AUROC ≥0.60 and naive−corrected <0.10. {hc.get('signatures_surviving_correction', 0)} passed; "
        f"{hc.get('signatures_surviving_correction_exploratory', 0)} passed looser exploratory criteria "
        "(≥0.55; Δ<0.15). The graveyard reports audit outcomes, not pathway discovery."
    )
    _fig(
        doc,
        fig_dir / "FigA3_signature_graveyard.png",
        "Naive pan-cohort versus within-indication corrected signature AUROC.",
        "Figure 5",
    )

    _h(doc, "Bulk RNA positive and null controls (E4)", 2)
    brca = tcga.get("per_indication", {}).get("TCGA-BRCA", tcga.get("primary_arm", {}))
    thca = tcga.get("per_indication", {}).get("TCGA-THCA", {})
    _p(
        doc,
        f"Balanced TCGA-BRCA bulk RNA (7 metastatic / 7 primary) yielded OOF AUROC {brca.get('oof_auroc', 0):.3f} "
        f"(primary positive control). Balanced TCGA-THCA (8+8) yielded OOF AUROC {thca.get('oof_auroc', 0):.3f} "
        "(null arm). TCGA-SKCM expression was degenerate (196/200 metastatic) and excluded."
    )
    _fig(
        doc,
        fig_dir / "FigA5_positive_control.png",
        "Bulk within-indication positive control versus confounded scRNA strata.",
        "Figure 6",
    )

    _h(doc, "Discussion", 1)
    _p(
        doc,
        "scClinicalAudit supports a cautious field law: on available public assets, apparent scRNA metastasis "
        "predictability is routinely bounded by metadata and provenance ceilings, while balanced bulk within-indication "
        "controls can still show detectable signal. The framework is intentionally conservative—its value is "
        "discriminating real from artifactual supervision, not maximizing AUROC."
    )
    _p(
        doc,
        "Limitations include pilot scale (342 registry-supervised specimens versus exhaustive atlas coverage), "
        "proxy signatures where gene-aligned scRNA indices are unavailable, partial portal holdout (transport only), "
        "and deferred F943 graph rebuild (~537 GB). Per-patient scRNA–genomics coupling awaits dbGaP barcodes."
    )

    _h(doc, "Methods", 1)
    _p(
        doc,
        "Software. scClinicalAudit v2.0 (`src/sc_clinical_audit/`). Headline estimands E1–E4 are defined in the "
        "OSF pre-registration protocol. Cross-validation used donor-grouped GroupKFold; pooled OOF AUROC is reported "
        "for ranking tasks. Specimen-authority labels reject bare disease-string metastasis positives."
    )
    _p(
        doc,
        "Data availability. Raw scRNA remains on HTAN and CZ CELLxGENE; bulk TCGA features via NCI GDC. Code, results, "
        "and figures are deposited on GitHub and Zenodo (Archetype A release). OSF pre-registration DOI to be inserted "
        "after upload."
    )
    _p(doc, "Competing interests. The authors declare no competing interests.")
    _p(doc, "Funding. This work received no specific external funding.")

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def build_supplementary(dr: Path, out: Path, fig_dir: Path) -> Path:
    report = _headline(dr)
    aa = dr / "artifacts/scientific_analyses/archetype_a"

    doc = Document()
    _style(doc)
    _h(doc, "Supplementary Information — scClinicalAudit Archetype A", 1)

    _h(doc, "Supplementary Methods", 2)
    _p(doc, "Estimand E1 (confound fraction): decompose_auroc_components() allocates AUROC above 0.5 proportionally to "
         "indication, portal, compositional, and provenance baselines up to the metadata ceiling.")
    _p(doc, "Estimand E2 (transportability): train block-logistic on stratum A, evaluate on stratum B.")
    _p(doc, "Estimand E3 (signature graveyard): strict survival requires corrected AUROC ≥0.60 and naive−corrected <0.10.")
    _p(doc, "Estimand E4 (positive control): balanced within-indication bulk RNA case-grouped CV; BRCA primary, THCA null.")

    if (aa / "specimen_registry.csv").is_file():
        reg = pd.read_csv(aa / "specimen_registry.csv")
        _table_from_df(doc, reg.head(30), "Supplementary Table S1 | Registry sample (first 30 rows).", 30)

    if (aa / "signature_graveyard.csv").is_file():
        sg = pd.read_csv(aa / "signature_graveyard.csv")
        cols = [c for c in sg.columns if c in (
            "signature_id", "name", "source", "naive_oof_auroc",
            "corrected_within_indication_auroc", "survives_correction",
            "survives_correction_exploratory",
        )]
        _table_from_df(doc, sg[cols], "Supplementary Table S2 | Signature graveyard.", 40)

    if (aa / "decomposition_by_pack.csv").is_file():
        _table_from_df(doc, pd.read_csv(aa / "decomposition_by_pack.csv"), "Supplementary Table S3 | Pack decomposition.", 20)

    if (aa / "transport_matrix_indication.csv").is_file():
        _table_from_df(doc, pd.read_csv(aa / "transport_matrix_indication.csv"), "Supplementary Table S4 | Transport matrix.", 15)

    _h(doc, "Supplementary Figures", 2)
    _fig(doc, fig_dir / "FigS1_portal_transport.png",
         "Portal-stratified transportability matrix.", "Supplementary Figure S1")

    doc.save(str(out))
    return out


def write_figure_legends(out_root: Path) -> Path:
    legends = out_root / "manuscript_figure_legends.md"
    text = """# Archetype A figure legends

**Figure 1.** Registry breadth and leakage taxonomy across C/D/E/F integrated manifests.

**Figure 2.** Confound decomposition spectrum: stacked AUROC above chance attributed to indication, portal, compositional, provenance, and residual components per pack.

**Figure 3.** Adversarial metadata ceilings versus TME block logistic regression on the canonical specimen-authority cohort.

**Figure 4.** Indication × indication transportability matrix (train stratum → test stratum AUROC).

**Figure 5.** Signature graveyard: naive pan-cohort versus within-indication corrected OOF AUROC for curated proxy programs.

**Figure 6.** Positive-control discriminator: balanced TCGA-BRCA bulk RNA versus pan-cohort and within-BRCA scRNA strata.

**Supplementary Figure S1.** Portal-stratified transportability on the canonical cohort.
"""
    legends.write_text(text, encoding="utf-8")
    return legends


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--data-root", type=Path, default=Path("D:/MetastaticEvolution"))
    ap.add_argument("--repo-root", type=Path, default=ROOT)
    ap.add_argument("--publish-root", type=Path, default=Path("C:/MetastaticEvolution/publish/archetype_a"))
    args = ap.parse_args()
    dr = args.data_root.resolve()
    pub = args.publish_root.resolve()
    pub.mkdir(parents=True, exist_ok=True)
    fig_dir = pub / "figures"
    out_dir = pub / "manuscript"
    out_dir.mkdir(parents=True, exist_ok=True)

    main_path = build_main(dr, out_dir / "Manuscript_ArchetypeA_Main.docx", fig_dir)
    supp_path = build_supplementary(dr, out_dir / "Manuscript_ArchetypeA_Supplementary.docx", fig_dir)
    legends = write_figure_legends(pub)

    repo_copy = args.repo_root / "Manuscript_ArchetypeA_Main.docx"
    shutil.copy2(main_path, repo_copy)

    print(f"Wrote {main_path}")
    print(f"Wrote {supp_path}")
    print(f"Wrote {legends}")
    print(f"Copied main manuscript to {repo_copy}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
